import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter as L
df=pd.read_pickle('all3.pkl')
df['time']=pd.to_datetime(df.time.astype(str),format='%H:%M:%S').dt.time
df['dow']=df.date.dt.day_name()
df=df.sort_values(['date','time']).reset_index(drop=True)
df['Doctor']=df.doc.map({'ND':'Dr. Dang','MI':'Dr. Iizuka'})
df['Month']=df.date.dt.strftime('%Y-%m')
YRS=['Sep 23–Aug 24','Sep 24–Aug 25','Sep 25–Sep 26']
df['Year']=pd.cut(df.date,[pd.Timestamp('2023-09-01'),pd.Timestamp('2024-09-01'),pd.Timestamp('2025-09-01'),pd.Timestamp('2026-10-01')],right=False,labels=YRS).astype(str)
CORR={'Macumira':'Follow Up <6 Mo'}
NONCLIN={'Surgery Cataract','Surgery Refractive','Surgery IOLX','MEETING'}
df['type_c']=df.type.replace(CORR)
ORDER=['Consult MD','Consult OD','Follow Up <6 Mo','Follow Up >6 Mo','LDD','LPI','Macumira','Post Op POM1 (Month)','Post Op POW1 (Week)','Post Op SD/POD1 (Day)','Scleral Lens Consult/Dispense','Scleral Lens Diagnostic','SLT','Surgery Counselling','Surgery IOLX','Surgery Cataract','Surgery Refractive','Virtual Call','YAG CAPS','YAG CAPS OD / OS']
EXTRA=[t for t in sorted(df.type.unique()) if t not in ORDER]
months=sorted(df.Month.unique());docs=['Dr. Dang','Dr. Iizuka']
F='Arial';H=Font(name=F,bold=True,color='FFFFFF');HF=PatternFill('solid',fgColor='1D5C8F')
B=Font(name=F,bold=True);N=Font(name=F);SUB=PatternFill('solid',fgColor='EEF2F5');GREY=Font(name=F,italic=True,color='5B6B78')
AMB=PatternFill('solid',fgColor='FBF0D9');WARN=Font(name=F,italic=True,color='9A6200')
wb=Workbook()
raw=wb.active;raw.title='Raw Data'
raw.append(['Demo #','Doctor','Appt Date','Appt Time','Weekday','Month','Type (corrected)','Reason','Roster','Type (as booked in Oscar)','Year (Sep–Aug)','Clinic visit?','Booked on schedule of'])
for r in df.itertuples():
    raw.append([r.demo,r.Doctor,r.date.date(),r.time,r.dow,r.Month,r.type_c,r.reason if isinstance(r.reason,str) else '',r.roster if isinstance(r.roster,str) else '',r.type,r.Year,'No' if r.type in NONCLIN else 'Yes',r.sched])
n=len(df)+1
for c in raw[1]: c.font=H;c.fill=HF
for row in raw.iter_rows(min_row=2):
    for c in row: c.font=N
    row[2].number_format='yyyy-mm-dd';row[3].number_format='h:mm'
    if row[6].value!=row[9].value:
        for c in row: c.fill=AMB
for i,w in enumerate([9,12,12,9,11,9,28,60,13,28,15,12,18],1): raw.column_dimensions[L(i)].width=w
raw.freeze_panes='A2';raw.auto_filter.ref=f'A1:M{n}'
R=lambda col:f"'Raw Data'!${col}$2:${col}${n}"
def header(ws,row,vals,start=1):
    for j,v in enumerate(vals,start):
        c=ws.cell(row,j,v);c.font=H;c.fill=HF;c.alignment=Alignment(horizontal='center' if j>1 else 'left',wrap_text=True,vertical='center')
def typerows(ws,start,ncols,fn):
    r=start
    for sect,types in [('Types you listed',ORDER),('Other types found in the export',EXTRA)]:
        ws.cell(r,1,sect).font=Font(name=F,bold=True,italic=True,color='5B6B78');r+=1
        for t in types:
            ws.cell(r,1,t).font=N
            for j in range(ncols): c=ws.cell(r,2+j,fn(r,j));c.font=N;c.number_format='#,##0;-#,##0;"-"'
            r+=1
    return r
def totals(ws,end,first,ncols,label='Total'):
    ws.cell(end,1,label).font=B;ws.cell(end,1).fill=SUB
    for j in range(ncols):
        col=L(2+j);c=ws.cell(end,2+j,f'=SUM({col}{first}:{col}{end-1})');c.font=B;c.fill=SUB;c.number_format='#,##0'
# Summary by type and year
s=wb.create_sheet('Summary by Type',0)
s['A1']='Appointment count by type, Sept 2023 – Sept 2026 (includes visits booked on Dr. Fan\'s and Dr. Tarn Gill\'s schedules)';s['A1'].font=Font(name=F,bold=True,size=13)
s['A2']='Each number counts rows on the Raw Data tab. Filter Raw Data by Doctor, Year and Type (corrected) to check any cell. Years run Sept–Aug; the last one also includes Sept 2026.';s['A2'].font=GREY
s['A4']='Before July 2025, Dr. Dang\'s and Dr. Iizuka\'s consults and follow-ups were booked on optometrist schedules. Dr. Fan\'s (Sept 2023 – Nov 2024) and Dr. Tarn Gill\'s (July 2024 – June 2025, rows tagged "ND"/"MI" in her Reason text) are now counted under the right doctor (Raw Data column M). Note: the July 2024 – June 2025 file was labelled "TG" and the Provider column reads "Gill, Tarn" throughout, not Dr. Schumacher as expected — flag this to Charoo if that\'s not right.';s['A4'].font=WARN
s['A3']='Correction: 1 Macumira visit (Dr. Dang, 2026-09-24) recounted as Follow Up <6 Mo. Highlighted on Raw Data; original type kept in column J.';s['A3'].font=WARN
s.cell(5,1,'');s.merge_cells('B5:E5');s.merge_cells('F5:I5')
s['B5']='Dr. Dang';s['F5']='Dr. Iizuka'
for a in ('B5','F5'): s[a].font=H;s[a].fill=HF;s[a].alignment=Alignment(horizontal='center')
s['J5'].fill=HF
header(s,6,['Appointment type']+YRS+['All years']+YRS+['All years','Both doctors, all years'])
def fs(r,j):
    if j in(0,1,2): return f'=COUNTIFS({R("G")},$A{r},{R("B")},"Dr. Dang",{R("K")},{L(2+j)}$6)'
    if j==3: return f'=SUM(B{r}:D{r})'
    if j in(4,5,6): return f'=COUNTIFS({R("G")},$A{r},{R("B")},"Dr. Iizuka",{R("K")},{L(2+j)}$6)'
    if j==7: return f'=SUM(F{r}:H{r})'
    return f'=E{r}+I{r}'
end=typerows(s,7,9,fs);totals(s,end,7,9,'Total rows')
s.cell(end+1,1,'Rows in Raw Data (should match)').font=GREY
c=s.cell(end+1,10,f'=COUNTA({R("A")})');c.font=GREY;c.number_format='#,##0'
s.column_dimensions['A'].width=40
for j in range(9): s.column_dimensions[L(2+j)].width=12.5
s.row_dimensions[6].height=32;s.freeze_panes='B7'
# Unique patients (values)
r=end+4
s.cell(r,1,'Unique patients (Demo #)').font=B;r+=1
header(s,r,['','Sep 23–Aug 24','Sep 24–Aug 25','Sep 25–Sep 26','All years']);r+=1
def uq(mask): return df[mask].demo.nunique()
for lab,m in [('Dr. Dang',df.Doctor=='Dr. Dang'),('Dr. Iizuka',df.Doctor=='Dr. Iizuka')]:
    s.cell(r,1,lab).font=N
    for j,y in enumerate(YRS): s.cell(r,2+j,uq(m&(df.Year==y))).number_format='#,##0'
    s.cell(r,5,uq(m)).number_format='#,##0';r+=1
s.cell(r,1,'Seen by both').font=N
for j,y in enumerate(YRS+[None]):
    sub=df if y is None else df[df.Year==y]
    s.cell(r,2+j,len(set(sub[sub.Doctor=='Dr. Dang'].demo)&set(sub[sub.Doctor=='Dr. Iizuka'].demo))).number_format='#,##0'
r+=1;s.cell(r,1,'Unique patients, practice').font=B
for j,y in enumerate(YRS+[None]):
    sub=df if y is None else df[df.Year==y];c=s.cell(r,2+j,sub.demo.nunique());c.font=B;c.number_format='#,##0'
r+=1;s.cell(r,1,'Counted from Raw Data in Python (Excel has no simple unique-count formula). Practice = Dr. Dang + Dr. Iizuka − seen by both.').font=GREY
# monthly per doctor
for dname,title in [('Dr. Dang','Dr. Dang by Month'),('Dr. Iizuka','Dr. Iizuka by Month')]:
    ws=wb.create_sheet(title)
    ws['A1']=f'{dname}: appointments by type and month, Sept 2023 – Sept 2026';ws['A1'].font=Font(name=F,bold=True,size=13)
    header(ws,3,['Appointment type']+months+['Total'])
    nm=len(months)
    def fm(r,j,dname=dname):
        if j<nm: return f'=COUNTIFS({R("G")},$A{r},{R("B")},"{dname}",{R("F")},{L(2+j)}$3)'
        return f'=SUM(B{r}:{L(1+nm)}{r})'
    end=typerows(ws,4,nm+1,fm);totals(ws,end,4,nm+1)
    ws.column_dimensions['A'].width=34
    for j in range(nm+1): ws.column_dimensions[L(2+j)].width=9
    ws.freeze_panes='B4'
# weekday
ws=wb.create_sheet('By Weekday')
ws['A1']='Appointments by type and weekday, all years';ws['A1'].font=Font(name=F,bold=True,size=13)
days=['Monday','Tuesday','Wednesday','Thursday','Friday','Saturday']
r0=3
for dname in docs:
    ws.cell(r0,1,dname).font=Font(name=F,bold=True,size=12);r0+=1
    header(ws,r0,['Appointment type']+days+['Total'])
    def fw(r,j,dname=dname):
        if j<6: return f'=COUNTIFS({R("G")},$A{r},{R("B")},"{dname}",{R("E")},"{days[j]}")'
        return f'=SUM(B{r}:G{r})'
    end=typerows(ws,r0+1,7,fw);totals(ws,end,r0+1,7)
    r0=end+3
ws.column_dimensions['A'].width=34
for j in range(7): ws.column_dimensions[L(2+j)].width=12
# same day
g=df.groupby(['demo','date','Doctor'])
pairs=g.agg(rows=('type_c','size'),types=('type_c',lambda x:' + '.join(sorted(x))),times=('time',lambda x:', '.join(t.strftime('%H:%M') for t in x)),Year=('Year','first')).reset_index()
pairs=pairs[pairs.rows>1].sort_values(['date','demo'])
ws=wb.create_sheet('Same-Day Rows')
ws['A1']='Patients with more than one appointment row on the same day';ws['A1'].font=Font(name=F,bold=True,size=13)
ws['A2']='One visit that shows as two rows: for example a consult with a same-day YAG or LPI, or bilateral cataract surgery (one row per eye).';ws['A2'].font=GREY
header(ws,4,['Demo #','Date','Doctor','Rows','Types','Times','Year'])
for i,p in enumerate(pairs.itertuples(),5):
    for j,v in enumerate([p.demo,p.date.date(),p.Doctor,p.rows,p.types,p.times,p.Year],1): ws.cell(i,j,v).font=N
    ws.cell(i,2).number_format='yyyy-mm-dd'
pe=4+len(pairs)
for i,w in enumerate([9,12,12,7,48,16,15],1): ws.column_dimensions[L(i)].width=w
ws.freeze_panes='A5';ws.auto_filter.ref=f'A4:G{pe}'
r=s.max_row+3
s.cell(r,1,'Same-day combinations (one visit, two rows), all years').font=B;r+=1
header(s,r,['Combination','Dr. Dang','Dr. Iizuka','Total']);r+=1
for t in pairs.types.value_counts().head(12).index:
    s.cell(r,1,t).font=N
    for j,dname in enumerate(docs): s.cell(r,2+j,f"=COUNTIFS('Same-Day Rows'!$E$5:$E${pe},$A{r},'Same-Day Rows'!$C$5:$C${pe},\"{dname}\")").font=N
    s.cell(r,4,f'=B{r}+C{r}').font=N;r+=1
s.cell(r,1,'All same-day patient-days').font=B
for j,dname in enumerate(docs): s.cell(r,2+j,f"=COUNTIFS('Same-Day Rows'!$C$5:$C${pe},\"{dname}\")").font=B
s.cell(r,4,f'=B{r}+C{r}').font=B
out='/mnt/user-data/outputs/Appointment_Type_Breakdown_ND_MI_2023-2026.xlsx';wb.save(out);print(out,len(pairs))

# ---- Monthly trend charts ----
from openpyxl.chart import LineChart,Reference
wb2=__import__('openpyxl').load_workbook(out)
rawn=n
ws=wb2.create_sheet('Monthly Trend',1)
ws['A1']='Monthly appointment volume, Sept 2023 – Sept 2026';ws['A1'].font=Font(name=F,bold=True,size=13)
ws['A2']='All appointment rows, including surgeries and same-day add-ons. Counts come from the Raw Data tab.';ws['A2'].font=GREY
ws['A3']='Dr. Dang and Dr. Iizuka include visits booked on Dr. Fan\'s schedule (Sept 2023 – Nov 2024) and Dr. Tarn Gill\'s schedule (July 2024 – June 2025, where her Reason text tagged the visit "ND" or "MI"). Dec 2024 – June 2025 still looks lower than the months around it — worth checking against OSCAR directly if that seems off.';ws['A3'].font=WARN
header(ws,4,['Month','Dr. Dang','Dr. Iizuka','Total'])
for i,m in enumerate(months,5):
    ws.cell(i,1,m).font=N
    ws.cell(i,2,f"=COUNTIFS('Raw Data'!$B$2:$B${rawn},\"Dr. Dang\",'Raw Data'!$F$2:$F${rawn},$A{i})").font=N
    ws.cell(i,3,f"=COUNTIFS('Raw Data'!$B$2:$B${rawn},\"Dr. Iizuka\",'Raw Data'!$F$2:$F${rawn},$A{i})").font=N
    ws.cell(i,4,f'=B{i}+C{i}').font=N
last=4+len(months)
for j,col in enumerate('BCD'):
    c=ws.cell(last+1,2+j,f'=SUM({col}5:{col}{last})');c.font=B;c.fill=SUB;c.number_format='#,##0'
ws.cell(last+1,1,'Total').font=B;ws.cell(last+1,1).fill=SUB
for c in 'ABCD': ws.column_dimensions[c].width=13
cats=Reference(ws,min_col=1,min_row=5,max_row=last)
def mk(title,cols,colors):
    ch=LineChart();ch.title=title;ch.height=9;ch.width=26;ch.y_axis.title='Appointments';ch.x_axis.title='Month'
    for col,colr in zip(cols,colors):
        ch.add_data(Reference(ws,min_col=col,min_row=4,max_row=last),titles_from_data=True)
        s_=ch.series[-1];s_.graphicalProperties.line.solidFill=colr;s_.graphicalProperties.line.width=22000;s_.smooth=False
    ch.set_categories(cats);ch.x_axis.tickLblSkip=3;ch.legend.position='b'
    ch.x_axis.delete=False;ch.y_axis.delete=False;ch.y_axis.majorGridlines=__import__('openpyxl').chart.axis.ChartLines()
    return ch
ws.add_chart(mk('Dr. Dang vs Dr. Iizuka, appointments per month',[2,3],['1D5C8F','C26A0A']),'F4')
ws.add_chart(mk('Total practice volume, appointments per month',[4],['2F6B5A']),'F24')

# ---- per clinic day ----
clin=df[~df.type.isin(NONCLIN)]
days=clin.groupby(['Doctor','Month']).date.nunique()
r0=4
hdr=['Dr. Dang clinic days','Dr. Dang clinic visits','Dr. Dang per clinic day','Dr. Iizuka clinic days','Dr. Iizuka clinic visits','Dr. Iizuka per clinic day']
for j,h in enumerate(hdr): 
    c=ws.cell(4,30+j,h);c.font=H;c.fill=HF;c.alignment=Alignment(wrap_text=True,horizontal='center')
ws.row_dimensions[4].height=45
for i,m in enumerate(months,5):
    for k,(dn,base) in enumerate([('Dr. Dang',30),('Dr. Iizuka',33)]):
        ws.cell(i,base,int(days.get((dn,m),0))).font=N
        ws.cell(i,base+1,f"=COUNTIFS('Raw Data'!$B$2:$B${rawn},\"{dn}\",'Raw Data'!$F$2:$F${rawn},$A{i},'Raw Data'!$L$2:$L${rawn},\"Yes\")").font=N
        c=ws.cell(i,base+2,f'=IFERROR({L(base+1)}{i}/{L(base)}{i},0)');c.font=N;c.number_format='0.0'
for j in range(30,36): ws.column_dimensions[L(j)].width=13
ws.cell(2,30,'Clinic days = dates with at least one clinic visit (counted in Python from Raw Data). Clinic visits exclude cataract, refractive and IOLX surgery and staff meetings (Raw Data column L).').font=GREY
ch=LineChart();ch.title='Clinic visits per clinic day, by month';ch.height=9;ch.width=26;ch.y_axis.title='Visits per clinic day';ch.x_axis.title='Month'
for col,colr in [(32,'1D5C8F'),(35,'C26A0A')]:
    ch.add_data(Reference(ws,min_col=col,min_row=4,max_row=last),titles_from_data=True)
    s_=ch.series[-1];s_.graphicalProperties.line.solidFill=colr;s_.graphicalProperties.line.width=22000
ch.set_categories(cats);ch.x_axis.tickLblSkip=3;ch.legend.position='b';ch.x_axis.delete=False;ch.y_axis.delete=False
ch.y_axis.majorGridlines=__import__('openpyxl').chart.axis.ChartLines()
ws.add_chart(ch,'F44')
print(days.unstack(0).to_string())

wb2.save(out);print('charts added')
